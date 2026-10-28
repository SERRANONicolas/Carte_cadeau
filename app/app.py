import io
import os
from datetime import date, datetime
from functools import wraps

from flask import (
    Flask, render_template, request, redirect, url_for,
    session, flash, send_file, abort
)
from sqlalchemy import inspect, text
from werkzeug.security import check_password_hash, generate_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from config import Config
from date_fr import date_fr_full
from models import (
    db, Demande, Enfant, AdminUser, Settings, SOCIETES, MAX_ENFANTS,
)


def create_app():
    app = Flask(__name__)
    app.config.from_object(Config)
    db.init_app(app)

    os.makedirs(Config.UPLOAD_DIR, exist_ok=True)

    with app.app_context():
        db.create_all()
        _ensure_admin_user()
        _ensure_settings()

    register_routes(app)

    @app.context_processor
    def inject_hero_image():
        settings = db.session.get(Settings, 1)
        classes = []
        style_parts = []
        if settings and settings.hero_image_filename:
            classes.append("has-photo")
            style_parts.append(f"--hero-bg-desktop:url('{url_for('media_hero_image')}')")
        if settings and settings.hero_image_mobile_filename:
            classes.append("has-mobile-photo")
            style_parts.append(f"--hero-bg-mobile:url('{url_for('media_hero_image_mobile')}')")
        return dict(
            hero_bg_class=" ".join(classes),
            hero_bg_style="; ".join(style_parts),
        )

    return app


def _ensure_admin_user():
    existing = AdminUser.query.filter_by(username=Config.ADMIN_USERNAME).first()
    if existing is None:
        admin = AdminUser(
            username=Config.ADMIN_USERNAME,
            password_hash=generate_password_hash(Config.ADMIN_PASSWORD),
        )
        db.session.add(admin)
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()


def _ensure_settings():
    # Migration légère : ajoute la colonne `annee` sur une base existante
    cols = [r[1] for r in db.session.execute(db.text("PRAGMA table_info(settings)"))]
    if cols and "annee" not in cols:
        db.session.execute(db.text(
            f"ALTER TABLE settings ADD COLUMN annee INTEGER NOT NULL DEFAULT {int(Config.DEFAULT_ANNEE)}"))
        db.session.commit()
    if db.session.get(Settings, 1) is None:
        db.session.add(Settings(
            id=1,
            annee=Config.DEFAULT_ANNEE,
            deadline_date=Config.DEFAULT_DEADLINE_DATE,
            enfants_nes_apres=Config.DEFAULT_ENFANTS_NES_APRES,
        ))
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("admin_id"):
            return redirect(url_for("admin_login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


def _is_closed(settings):
    return settings.deadline_date is not None and date.today() > settings.deadline_date


def _public_context(settings, form=None):
    return dict(
        annee=settings.annee,
        societes=SOCIETES,
        rangs=range(1, MAX_ENFANTS + 1),
        enfants_nes_apres=date_fr_numeric(settings.enfants_nes_apres),
        deadline_full=date_fr_full(settings.deadline_date),
        is_closed=_is_closed(settings),
        form=form,
        show_snow=False,
    )


def date_fr_numeric(d):
    return d.strftime("%d/%m/%Y")


def register_routes(app):

    @app.get("/")
    def index():
        settings = db.session.get(Settings, 1)
        return render_template("index.html", **_public_context(settings))

    @app.post("/demande")
    def demande():
        settings = db.session.get(Settings, 1)

        if _is_closed(settings):
            flash("Les demandes sont closes depuis le "
                  f"{date_fr_full(settings.deadline_date)}.", "error")
            return render_template("index.html", **_public_context(settings)), 403

        societe = (request.form.get("societe") or "").strip()
        nom = (request.form.get("nom_salarie") or "").strip()
        prenom = (request.form.get("prenom_salarie") or "").strip()
        enfant_commun = request.form.get("enfant_commun") == "on"

        errors = []
        if societe not in SOCIETES:
            errors.append("Merci de sélectionner votre société.")
        if not nom:
            errors.append("Le nom du salarié est obligatoire.")
        if not prenom:
            errors.append("Le prénom du salarié est obligatoire.")

        enfants = []
        for rang in range(1, MAX_ENFANTS + 1):
            nom_prenom = (request.form.get(f"enfant_{rang}") or "").strip()
            if nom_prenom:
                enfants.append((rang, nom_prenom))

        if not enfants:
            errors.append("Merci de renseigner au moins un enfant.")

        if errors:
            for e in errors:
                flash(e, "error")
            return render_template(
                "index.html", **_public_context(settings, form=request.form)
            ), 400

        d = Demande(
            societe=societe, nom_salarie=nom, prenom_salarie=prenom,
            enfant_commun=enfant_commun,
        )
        db.session.add(d)
        db.session.flush()
        for rang, nom_prenom in enfants:
            db.session.add(Enfant(demande_id=d.id, rang=rang, nom_prenom=nom_prenom))
        db.session.commit()

        return render_template("confirmation.html", prenom_salarie=prenom)

    # ---------------- Admin ----------------

    @app.get("/admin")
    def admin_root():
        return redirect(url_for("admin_dashboard"))

    @app.route("/admin/login", methods=["GET", "POST"])
    def admin_login():
        if request.method == "POST":
            username = (request.form.get("username") or "").strip()
            password = request.form.get("password") or ""
            admin = AdminUser.query.filter_by(username=username).first()
            if admin and check_password_hash(admin.password_hash, password):
                session.clear()
                session["admin_id"] = admin.id
                session.permanent = True
                next_url = request.args.get("next") or url_for("admin_dashboard")
                return redirect(next_url)
            flash("Identifiants incorrects.", "error")
        return render_template("admin_login.html", show_snow=False)

    @app.get("/admin/logout")
    def admin_logout():
        session.clear()
        return redirect(url_for("admin_login"))

    @app.get("/admin/dashboard")
    @login_required
    def admin_dashboard():
        demandes = Demande.query.order_by(Demande.date_creation.desc()).all()
        par_societe = {}
        for d in demandes:
            par_societe[d.societe] = par_societe.get(d.societe, 0) + 1
        return render_template(
            "admin_dashboard.html",
            demandes=demandes,
            total_demandes=len(demandes),
            total_enfants=sum(len(d.enfants) for d in demandes),
            total_communs=sum(1 for d in demandes if d.enfant_commun),
            par_societe=par_societe,
        )

    @app.get("/admin/demandes/<int:demande_id>/delete")
    @login_required
    def admin_delete(demande_id):
        d = db.get_or_404(Demande, demande_id)
        db.session.delete(d)
        db.session.commit()
        flash("Demande supprimée.", "success")
        return redirect(url_for("admin_dashboard"))

    @app.post("/admin/reset-demandes")
    @login_required
    def admin_reset_demandes():
        confirm_text = (request.form.get("confirm") or "").strip().upper()
        if confirm_text != "RAZ":
            flash("Réinitialisation annulée : le mot de confirmation ne correspond pas.", "error")
            return redirect(url_for("admin_settings"))

        count = Demande.query.count()
        Enfant.query.delete()
        Demande.query.delete()
        db.session.commit()

        flash(f"{count} demande(s) supprimée(s). La base est réinitialisée.", "success")
        return redirect(url_for("admin_settings"))

    @app.get("/admin/export.xlsx")
    @login_required
    def admin_export():
        demandes = Demande.query.order_by(Demande.date_creation.asc()).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Demandes"

        headers = ["Date de la demande", "Société", "Nom", "Prénom", "Nb enfants"]
        headers += [f"Enfant {i}" for i in range(1, MAX_ENFANTS + 1)]
        headers += ["Enfant en commun avec un autre salarié"]
        header_fill = PatternFill(start_color="1F3B2C", end_color="1F3B2C", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        for d in demandes:
            par_rang = {e.rang: e.nom_prenom for e in d.enfants}
            ws.append(
                [d.date_creation.strftime("%d/%m/%Y %H:%M"), d.societe,
                 d.nom_salarie, d.prenom_salarie, len(d.enfants)]
                + [par_rang.get(i, "") for i in range(1, MAX_ENFANTS + 1)]
                + ["Oui" if d.enfant_commun else ""]
            )

        widths = [18, 30, 20, 20, 11] + [26] * MAX_ENFANTS + [22]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f"demandes_cartes_cadeaux_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return send_file(
            buf,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # ---------------- Média ----------------

    def _serve_settings_image(field_name):
        settings = db.session.get(Settings, 1)
        filename = getattr(settings, field_name, None) if settings else None
        if not filename:
            abort(404)
        path = os.path.join(Config.UPLOAD_DIR, filename)
        if not os.path.isfile(path):
            abort(404)
        return send_file(path)

    @app.get("/media/hero-image")
    def media_hero_image():
        return _serve_settings_image("hero_image_filename")

    @app.get("/media/hero-image-mobile")
    def media_hero_image_mobile():
        return _serve_settings_image("hero_image_mobile_filename")

    # ---------------- Admin : paramètres ----------------

    # ---------------- Admin : paramètres ----------------

    @app.get("/admin/settings")
    @login_required
    def admin_settings():
        settings = db.session.get(Settings, 1)
        return render_template(
            "admin_settings.html",
            settings=settings,
            hero_image_ideal_resolution=Config.HERO_IMAGE_IDEAL_RESOLUTION,
            hero_image_mobile_ideal_resolution=Config.HERO_IMAGE_MOBILE_IDEAL_RESOLUTION,
            total_inscriptions=Demande.query.count(),
        )

    @app.post("/admin/settings/general")
    @login_required
    def admin_settings_general():
        settings = db.session.get(Settings, 1)

        nes_apres_raw = request.form.get("enfants_nes_apres") or ""
        deadline_raw = request.form.get("deadline_date") or ""

        errors = []
        try:
            annee = int(request.form.get("annee") or "")
            if not 2000 <= annee <= 2100:
                raise ValueError
            settings.annee = annee
        except ValueError:
            errors.append("Année invalide.")
        try:
            settings.enfants_nes_apres = date.fromisoformat(nes_apres_raw)
        except ValueError:
            errors.append("Date de naissance minimale invalide.")

        if deadline_raw.strip():
            try:
                settings.deadline_date = date.fromisoformat(deadline_raw)
            except ValueError:
                errors.append("Date limite de demande invalide.")
        else:
            settings.deadline_date = None

        if errors:
            for e in errors:
                flash(e, "error")
            db.session.rollback()
        else:
            db.session.commit()
            flash("Informations générales mises à jour.", "success")

        return redirect(url_for("admin_settings"))

    def _handle_hero_image_upload(field_name, upload_prefix, success_label):
        settings = db.session.get(Settings, 1)
        current_filename = getattr(settings, field_name)

        if request.form.get("action") == "remove":
            if current_filename:
                try:
                    os.remove(os.path.join(Config.UPLOAD_DIR, current_filename))
                except OSError:
                    pass
                setattr(settings, field_name, None)
                db.session.commit()
            flash(f"{success_label} supprimée.", "success")
            return redirect(url_for("admin_settings"))

        file = request.files.get("image")
        if not file or not file.filename:
            flash("Merci de sélectionner une image.", "error")
            return redirect(url_for("admin_settings"))

        ext = os.path.splitext(file.filename)[1].lower()
        if ext not in Config.HERO_IMAGE_ALLOWED_EXT:
            flash("Format non supporté. Utilisez JPG, PNG ou WEBP.", "error")
            return redirect(url_for("admin_settings"))

        file.seek(0, os.SEEK_END)
        size = file.tell()
        file.seek(0)
        if size > Config.HERO_IMAGE_MAX_BYTES:
            flash("Fichier trop volumineux (8 Mo maximum).", "error")
            return redirect(url_for("admin_settings"))

        if current_filename:
            try:
                os.remove(os.path.join(Config.UPLOAD_DIR, current_filename))
            except OSError:
                pass

        filename = f"{upload_prefix}{ext}"
        file.save(os.path.join(Config.UPLOAD_DIR, filename))
        setattr(settings, field_name, filename)
        db.session.commit()

        flash(f"{success_label} mise à jour.", "success")
        return redirect(url_for("admin_settings"))

    @app.post("/admin/settings/hero-image")
    @login_required
    def admin_settings_hero_image():
        return _handle_hero_image_upload("hero_image_filename", "hero", "Image de fond (desktop)")

    @app.post("/admin/settings/hero-image-mobile")
    @login_required
    def admin_settings_hero_image_mobile():
        return _handle_hero_image_upload("hero_image_mobile_filename", "hero-mobile", "Image de fond (mobile)")

    @app.errorhandler(404)
    def not_found(e):
        return render_template("404.html", show_snow=False), 404


app = create_app()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)
